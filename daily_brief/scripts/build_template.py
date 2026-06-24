"""Builds templates/brief_template.docx: the styled base document that
generate_brief.py fills in. Run once (or whenever the look needs to change):

    python scripts/build_template.py
"""

import os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor, Inches

TEMPLATE_PATH = os.path.join(os.path.dirname(__file__), "..", "templates", "brief_template.docx")

ACCENT_COLOR = RGBColor(0x1F, 0x3A, 0x5F)
MUTED_COLOR = RGBColor(0x5A, 0x5A, 0x5A)


def build_template() -> None:
    doc = Document()

    for section in doc.sections:
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        section.top_margin = Inches(0.9)
        section.bottom_margin = Inches(0.9)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    title = doc.styles["Title"]
    title.font.name = "Calibri"
    title.font.size = Pt(28)
    title.font.bold = True
    title.font.color.rgb = ACCENT_COLOR

    subtitle = doc.styles["Subtitle"]
    subtitle.font.name = "Calibri"
    subtitle.font.size = Pt(14)
    subtitle.font.italic = True
    subtitle.font.color.rgb = MUTED_COLOR

    heading1 = doc.styles["Heading 1"]
    heading1.font.name = "Calibri"
    heading1.font.size = Pt(16)
    heading1.font.bold = True
    heading1.font.color.rgb = ACCENT_COLOR

    quote_style = doc.styles["Quote"]
    quote_style.font.name = "Calibri"
    quote_style.font.size = Pt(13)
    quote_style.font.italic = True
    quote_style.paragraph_format.left_indent = Inches(0.4)
    quote_style.paragraph_format.space_before = Pt(12)
    quote_style.paragraph_format.space_after = Pt(2)

    explanation_style = doc.styles.add_style("Quote Explanation", normal.type)
    explanation_style.base_style = normal
    explanation_style.font.size = Pt(10.5)
    explanation_style.font.color.rgb = MUTED_COLOR
    explanation_style.paragraph_format.left_indent = Inches(0.4)
    explanation_style.paragraph_format.space_after = Pt(10)

    placeholder_title = doc.add_paragraph("Person Name", style="Title")
    placeholder_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    placeholder_subtitle = doc.add_paragraph("Short Description", style="Subtitle")
    placeholder_subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

    os.makedirs(os.path.dirname(TEMPLATE_PATH), exist_ok=True)
    doc.save(TEMPLATE_PATH)
    print(f"Template written to {TEMPLATE_PATH}")


if __name__ == "__main__":
    build_template()
