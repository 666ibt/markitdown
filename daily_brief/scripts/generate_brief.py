"""Generates today's brief: picks a random person from data/people.json,
researches them with Claude (web search, no Wikipedia, books/articles
preferred), and writes a 2-3 page .docx brief into output/<date>/.

Usage:
    export ANTHROPIC_API_KEY=...
    python scripts/generate_brief.py
"""

import json
import os
import random
import re
from datetime import date

import anthropic
from docx import Document

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
PROJECT_DIR = os.path.dirname(SCRIPT_DIR)
PEOPLE_PATH = os.path.join(PROJECT_DIR, "data", "people.json")
TEMPLATE_PATH = os.path.join(PROJECT_DIR, "templates", "brief_template.docx")
OUTPUT_DIR = os.path.join(PROJECT_DIR, "output")

MODEL = "claude-opus-4-8"
BLOCKED_DOMAINS = ["wikipedia.org", "en.wikipedia.org", "simple.wikipedia.org"]

SYSTEM_PROMPT = """You are a research assistant producing a concise, well-sourced \
biographical brief for a personal daily-reading project. You write in clear, \
engaging prose suitable for an educated general reader."""

USER_PROMPT_TEMPLATE = """Research {name} ({category}) using the web_search tool.

Source rules:
- Do NOT use wikipedia.org (or any Wikipedia mirror) as a source.
- Prioritize books, biographies, essays, long-form journalism, and reputable \
articles over generic web pages.
- Cross-check key facts and quotes across at least two independent sources \
when possible.

IMPORTANT: All text values in your response must be written in RUSSIAN \
(person_name, short_description, intro, key_points, quote explanations, and \
the quotes themselves translated into Russian where the original is not Russian).

After researching, respond with ONLY a single JSON object - no markdown code \
fences, no commentary before or after - matching exactly this schema:

{{
  "person_name": "Полное имя, под которым человек широко известен (на русском)",
  "short_description": "Краткое описание из 1-3 слов для имени файла, напр. 'Философ-стоик'",
  "lifespan": "напр. 121-180 гг. н.э., или пустая строка, если неприменимо",
  "intro": "Биографическое введение из 3-5 абзацев (кто это, когда и где жил, \
почему важен), обычным текстом, абзацы разделены пустой строкой. На русском языке.",
  "key_points": ["5-8 строк, каждая - крупное достижение или ключевая идея из его работы. На русском языке."],
  "quotes": [
    {{"quote": "Подлинная цитата, приписываемая этому человеку (на русском)", \
"explanation": "1-3 предложения, объясняющие смысл цитаты и почему она важна. На русском языке."}}
  ]
}}

Include between 3 and 5 entries in "quotes". Make sure every quote is genuinely \
attributable to {name} based on your research. Remember: every text value MUST be in Russian."""


def load_people() -> list[dict]:
    with open(PEOPLE_PATH, "r", encoding="utf-8") as f:
        return json.load(f)


def save_people(people: list[dict]) -> None:
    with open(PEOPLE_PATH, "w", encoding="utf-8") as f:
        json.dump(people, f, ensure_ascii=False, indent=2)
        f.write("\n")


def pick_person(people: list[dict]) -> dict:
    return random.choice(people)


def remove_person(people: list[dict], person: dict) -> None:
    """Drop the picked person from the list and persist it, so they are
    never selected again."""
    people.remove(person)
    save_people(people)


def extract_json(text: str) -> dict:
    text = text.strip()
    fenced = re.search(r"```(?:json)?\s*(.*?)```", text, re.DOTALL)
    if fenced:
        text = fenced.group(1).strip()
    start = text.find("{")
    end = text.rfind("}")
    if start == -1 or end == -1:
        raise ValueError(f"No JSON object found in model response:\n{text}")
    return json.loads(text[start : end + 1])


def research_person(client: anthropic.Anthropic, person: dict) -> dict:
    user_prompt = USER_PROMPT_TEMPLATE.format(
        name=person["name"], category=person["category"]
    )

    with client.messages.stream(
        model=MODEL,
        max_tokens=16000,
        thinking={"type": "adaptive"},
        system=SYSTEM_PROMPT,
        tools=[
            {
                "type": "web_search_20260209",
                "name": "web_search",
                "max_uses": 8,
                "blocked_domains": BLOCKED_DOMAINS,
            }
        ],
        messages=[{"role": "user", "content": user_prompt}],
    ) as stream:
        for _ in stream.text_stream:
            pass
        final_message = stream.get_final_message()

    text_blocks = [block.text for block in final_message.content if block.type == "text"]
    if not text_blocks:
        raise RuntimeError(f"No text content returned. stop_reason={final_message.stop_reason}")

    return extract_json(text_blocks[-1])


def sanitize_filename(text: str) -> str:
    text = re.sub(r'[\\/:*?"<>|]', "", text)
    return text.strip()


def build_document(data: dict) -> Document:
    doc = Document(TEMPLATE_PATH)

    doc.paragraphs[0].runs[0].text = data["person_name"]
    subtitle_text = data["short_description"]
    if data.get("lifespan"):
        subtitle_text += f"  ({data['lifespan']})"
    doc.paragraphs[1].runs[0].text = subtitle_text

    doc.add_heading("Кто это", level=1)
    for paragraph in data["intro"].split("\n\n"):
        paragraph = paragraph.strip()
        if paragraph:
            doc.add_paragraph(paragraph)

    doc.add_page_break()
    doc.add_heading("Достижения и ключевые идеи", level=1)
    for point in data["key_points"]:
        doc.add_paragraph(point, style="List Bullet")

    doc.add_page_break()
    doc.add_heading("Цитаты", level=1)
    for item in data["quotes"]:
        doc.add_paragraph(f"“{item['quote']}”", style="Quote")
        doc.add_paragraph(item["explanation"], style="Quote Explanation")

    return doc


def save_document(doc: Document, data: dict) -> str:
    today_dir = os.path.join(OUTPUT_DIR, date.today().isoformat())
    os.makedirs(today_dir, exist_ok=True)

    name = sanitize_filename(data["person_name"])
    description = sanitize_filename(data["short_description"])
    filename = f"{name} - {description}.docx"
    path = os.path.join(today_dir, filename)

    doc.save(path)
    return path


def main() -> None:
    people = load_people()
    person = pick_person(people)
    print(f"Selected: {person['name']} ({person['category']})")

    client = anthropic.Anthropic()
    print("Researching...")
    data = research_person(client, person)

    print("Building document...")
    doc = build_document(data)
    path = save_document(doc, data)

    remove_person(people, person)

    print(f"Saved brief to: {path}")
    print(f"Removed {person['name']} from the list ({len(people)} remaining).")


if __name__ == "__main__":
    main()
